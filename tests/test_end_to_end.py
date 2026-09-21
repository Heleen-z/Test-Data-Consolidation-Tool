from __future__ import annotations

from openpyxl import load_workbook
import pandas as pd

from testdata_tool.config import load_config
from testdata_tool.models import FileStatus
from testdata_tool.pipeline import merge_prechecked
from testdata_tool.precheck import run_precheck


def test_precheck_and_multi_type_export(sample_workspace):
    input_folder, output_folder, template_folder = sample_workspace
    config = load_config()
    precheck = run_precheck(input_folder, template_folder, output_folder, config)

    assert precheck.scanned_count == 3
    assert precheck.valid_count == 3
    assert {item.test_type for item in precheck.files} == {"FAKT", "GANTER"}
    assert all(item.status in (FileStatus.VALID, FileStatus.WARNING) for item in precheck.files)

    result = merge_prechecked(precheck, output_folder, config)
    assert result.output_path.exists()
    assert result.type_rows == {"FAKT": 4, "GANTER": 2}

    fakt = pd.read_parquet(result.datasets["FAKT"])
    assert list(fakt["__source_row"]) == [2, 3, 2, 3]
    assert "Unused" not in fakt.columns
    assert list(fakt["Runtime [ms]"]) == [0, 100, 0, 100]

    workbook = load_workbook(result.output_path, read_only=False, data_only=True)
    assert {"Summary", "FAKT", "GANTER", "Source_Index", "Validation"} <= set(workbook.sheetnames)
    summary = {
        row[0].value: row[1].value
        for row in workbook["Summary"].iter_rows(min_row=2, max_col=2)
        if row[0].value
    }
    assert summary["程序版本"] == "1.1.1"
    assert summary["配置版本"] == "built-in"
    assert summary["模板版本"] == "built-in"
    sheet = workbook["FAKT"]
    headers = [cell.value for cell in sheet[1]]
    assert "Unused" not in headers
    assert sheet.column_dimensions["D"].hidden or any(
        sheet.column_dimensions[cell.column_letter].hidden
        for cell in sheet[1]
        if str(cell.value).startswith("__")
    )
    workbook.close()


def test_merge_can_include_only_user_selected_valid_files(sample_workspace):
    input_folder, output_folder, template_folder = sample_workspace
    config = load_config()
    precheck = run_precheck(input_folder, template_folder, output_folder, config)
    chosen = next(item for item in precheck.files if item.path.name == "fakt_a.txt")

    result = merge_prechecked(
        precheck, output_folder, config, selected_source_ids={chosen.source_id}
    )

    assert result.type_rows == {"FAKT": 2}
    assert set(result.source_index["Source_ID"]) == {chosen.source_id}
    assert result.validation.loc[
        result.validation["source_id"] == chosen.source_id, "是否选择合并"
    ].item()
    assert not result.validation.loc[
        result.validation["source_id"] != chosen.source_id, "是否选择合并"
    ].any()


def test_exact_duplicate_is_skipped(sample_workspace):
    input_folder, output_folder, template_folder = sample_workspace
    duplicate = input_folder / "copy.txt"
    duplicate.write_bytes((input_folder / "fakt_a.txt").read_bytes())
    precheck = run_precheck(input_folder, template_folder, output_folder, load_config())
    assert sum(item.status == FileStatus.DUPLICATE for item in precheck.files) == 1


def test_missing_axis_is_invalid(sample_workspace):
    input_folder, output_folder, template_folder = sample_workspace
    (input_folder / "fakt_no_axis.txt").write_text(
        "Temperature actual value [C];Unused\n20;1\n", encoding="utf-8"
    )
    precheck = run_precheck(input_folder, template_folder, output_folder, load_config())
    item = next(item for item in precheck.files if item.path.name == "fakt_no_axis.txt")
    assert item.status == FileStatus.INVALID
    assert any("横轴" in message for message in item.messages)
