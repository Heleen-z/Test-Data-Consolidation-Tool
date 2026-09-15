from pathlib import Path

from openpyxl import load_workbook

from testdata_tool.config import load_config
from testdata_tool.templates import load_templates


def test_supplied_templates_are_detected():
    config = load_config()
    templates = load_templates(Path("Raw data standard Title"), config)
    assert set(templates) == {
        "FAKT",
        "GANTER",
        "MP_SHAKER",
        "SEOUL",
        "T09",
        "T12",
        "T13",
    }
    assert "TCK01" in templates["GANTER"].fields
    assert "X-Data [s]" in templates["MP_SHAKER"].fields
    assert all(template.fields for template in templates.values())


def test_supplied_templates_do_not_contain_degree_sign_lookalikes():
    lookalikes = {"ｰ", "ー", "º", "˚"}
    findings = []
    for path in Path("Raw data standard Title").glob("*.xlsx"):
        workbook = load_workbook(path, read_only=True, data_only=False)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(
                        char in cell.value for char in lookalikes
                    ):
                        findings.append(f"{path.name}:{sheet.title}!{cell.coordinate}")
        workbook.close()
    assert findings == []
