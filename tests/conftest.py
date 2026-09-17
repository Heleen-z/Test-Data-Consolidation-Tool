from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


YELLOW = PatternFill("solid", fgColor="FFFF00")
GREEN = PatternFill("solid", fgColor="00B050")


@pytest.fixture(autouse=True)
def isolated_usage_network(tmp_path, monkeypatch):
    """All tests use an isolated queue; production transport is never a test fixture."""
    from testdata_tool.usage_logging import SharePointClient, UploadError, _services
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "profile"))
    from testdata_tool.usage_logging import UsageService
    original_wake = UsageService.wake
    monkeypatch.setattr(UsageService, "wake", lambda self: None)
    def blocked(*args, **kwargs):
        raise UploadError("自动测试禁止访问真实 SharePoint")
    monkeypatch.setattr(SharePointClient, "__init__", _isolated_client_init(SharePointClient.__init__, blocked))
    yield original_wake
    for service in list(_services.values()):
        service._close_bridge()
    _services.clear()


def _isolated_client_init(original, blocked):
    def init(self, settings, request=None):
        original(self, settings, request=blocked if request is None else request)
    return init


def make_template(path: Path, headers: list[str], selected: set[str], header_row: int = 1):
    workbook = Workbook()
    sheet = workbook.active
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, value)
        if value in selected:
            cell.fill = GREEN if column % 2 == 0 else YELLOW
    workbook.save(path)


@pytest.fixture
def sample_workspace(tmp_path):
    input_folder = tmp_path / "input"
    output_folder = tmp_path / "output"
    template_folder = tmp_path / "templates"
    input_folder.mkdir()
    output_folder.mkdir()
    template_folder.mkdir()

    make_template(
        template_folder / "FAKT 原始记录文件表头.xlsx",
        ["Timestamp [ms]", "Runtime [ms]", "Temperature actual value [C]", "Unused"],
        {"Runtime [ms]", "Temperature actual value [C]"},
    )
    make_template(
        template_folder / "GANTER 原始记录文件表头.xlsx",
        ["Date", "Time", "TCK01", "TCK02", "Unused"],
        {"TCK01", "TCK02"},
    )
    (input_folder / "fakt_a.txt").write_text(
        "Timestamp [ms];Runtime [ms];Temperature actual value [C];Unused\n"
        "0;0;20.5;999\n100;100;21.0;999\n",
        encoding="utf-8",
    )
    (input_folder / "fakt_b.txt").write_text(
        "Timestamp [ms];Runtime [ms];Temperature actual value [C]\n"
        "0;0;22.0\n100;100;23.0\n",
        encoding="utf-8",
    )
    (input_folder / "ganter.csv").write_text(
        "Date,Time,TCK01,TCK02,Other\n"
        "2026-08-01,10:00:00,100,110,1\n"
        "2026-08-01,10:00:01,101,111,2\n",
        encoding="utf-8",
    )
    return input_folder, output_folder, template_folder
