from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterator
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from ..models import FileInspection, HeaderInspection
from ..normalize import normalize_header


TEXT_EXTENSIONS = {".csv", ".txt"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


def detect_encoding(path: Path) -> str:
    with path.open("rb") as handle:
        raw = handle.read(128 * 1024)

    # A German/European device export containing ``\xb0C`` is also valid GB18030:
    # GB18030 interprets both bytes as one CJK character and silently drops the C.
    # Therefore a first-success strategy is unsafe.  Compare every strict decode
    # instead.  Correct decodes normally need the fewest non-ASCII characters;
    # CJK is only a tie-breaker, so genuine Chinese exports still select GB18030.
    candidates: list[tuple[tuple[int, int, int, int], int, str]] = []
    preference = ("utf-8-sig", "utf-8", "gb18030", "cp1252")
    for order, encoding in enumerate(preference):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        non_ascii = sum(ord(char) > 127 for char in text)
        controls = sum(
            unicodedata.category(char) == "Cc" and char not in "\r\n\t"
            for char in text
        )
        cjk = sum("\u3400" <= char <= "\u9fff" for char in text)
        # A byte stream such as B0 43 decodes as one CJK character in GB18030,
        # but as degree sign + ASCII C in CP1252.  Prefer the latter only when
        # the primary non-ASCII score is tied.
        candidates.append(((controls, non_ascii, cjk, order), order, encoding))
    if candidates:
        return min(candidates)[2]
    return "latin1"


def detect_delimiter(lines: list[str]) -> str:
    sample = "\n".join(lines[:10])
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        scores = {delimiter: sample.count(delimiter) for delimiter in (";", ",", "\t", "|")}
        return max(scores, key=scores.get)


def _score_row(values: list[object], expected: set[str], axes: set[str]) -> int:
    normalized = {normalize_header(value) for value in values if value is not None}
    return 100 * len(normalized & expected) + 30 * len(normalized & axes) + min(20, len(normalized))


class TabularParser:
    parser_id = "tabular"

    def probe(self, path: Path) -> bool:
        return path.suffix.casefold() in TEXT_EXTENSIONS | EXCEL_EXTENSIONS

    def inspect(
        self, path: Path, expected_fields: list[str], axis_candidates: list[str]
    ) -> HeaderInspection:
        expected = {normalize_header(item) for item in expected_fields}
        axes = {normalize_header(item) for item in axis_candidates}
        if path.suffix.casefold() in EXCEL_EXTENSIONS:
            workbook = load_workbook(
                path, read_only=True, data_only=True, keep_links=False
            )
            sheet = workbook.active
            preview = [list(row) for row in sheet.iter_rows(max_row=30, values_only=True)]
            row_count = max(0, sheet.max_row)
            workbook.close()
            if not preview:
                raise ValueError("工作表为空")
            scores = [_score_row(row, expected, axes) for row in preview]
            header_index = max(range(len(scores)), key=scores.__getitem__)
            columns = ["" if value is None else str(value).strip() for value in preview[header_index]]
            while columns and not columns[-1]:
                columns.pop()
            return HeaderInspection(columns, header_index, max(0, row_count - header_index - 1))

        encoding = detect_encoding(path)
        with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
            lines = [handle.readline() for _ in range(40)]
        lines = [line for line in lines if line]
        if not lines:
            raise ValueError("文本文件为空")
        delimiter = detect_delimiter(lines)
        rows = list(csv.reader(lines, delimiter=delimiter))
        scores = [_score_row(row, expected, axes) for row in rows]
        header_index = max(range(len(scores)), key=scores.__getitem__)
        columns = [str(value).strip() for value in rows[header_index]]
        with path.open("r", encoding=encoding, errors="replace") as handle:
            total_lines = sum(1 for _ in handle)
        return HeaderInspection(
            columns,
            header_index,
            max(0, total_lines - header_index - 1),
            delimiter,
            encoding,
        )

    def iter_batches(
        self, path: Path, inspection: FileInspection, chunk_rows: int
    ) -> Iterator[pd.DataFrame]:
        if path.suffix.casefold() in EXCEL_EXTENSIONS:
            frame = pd.read_excel(path, header=inspection.header_row, engine="openpyxl")
            frame.columns = [str(item).strip() for item in frame.columns]
            for start in range(0, len(frame), chunk_rows):
                yield frame.iloc[start : start + chunk_rows].copy()
            return
        reader = pd.read_csv(
            path,
            header=inspection.header_row,
            sep=inspection.delimiter or None,
            encoding=inspection.encoding or "utf-8-sig",
            engine="python",
            chunksize=chunk_rows,
            on_bad_lines="warn",
        )
        for frame in reader:
            frame.columns = [str(item).strip() for item in frame.columns]
            yield frame


class DelimitedParser(TabularParser):
    parser_id = "delimited"


class UnsupportedParser:
    def __init__(self, parser_id: str):
        self.parser_id = parser_id

    def probe(self, path: Path) -> bool:
        return False

    def inspect(self, path: Path, expected_fields: list[str], axis_candidates: list[str]):
        raise ValueError(f"{self.parser_id} 属于第二阶段专用格式，当前版本尚未支持")

    def iter_batches(self, path: Path, inspection: FileInspection, chunk_rows: int):
        raise ValueError(f"{self.parser_id} 属于第二阶段专用格式，当前版本尚未支持")
