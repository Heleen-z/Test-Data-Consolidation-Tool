from __future__ import annotations

import warnings
from pathlib import Path

from openpyxl import load_workbook

from .config import AppConfig
from .models import TemplateDefinition
from .normalize import normalize_header


class TemplateError(ValueError):
    pass


def _cell_rgb(cell) -> str:
    fill = getattr(cell, "fill", None)
    if fill is None or getattr(fill, "fill_type", None) != "solid":
        return ""
    color = fill.fgColor
    if color.type == "rgb" and color.rgb:
        return str(color.rgb).upper()
    return ""


def _rgb_key(value: str) -> str:
    value = str(value).upper().replace("#", "")
    return value[-6:] if len(value) >= 6 else value


def _template_type(path: Path, keywords: dict[str, list[str]]) -> str | None:
    name = normalize_header(path.stem)
    matches = [
        test_type
        for test_type, words in keywords.items()
        if any(normalize_header(word) in name for word in words)
    ]
    return matches[0] if len(matches) == 1 else None


def load_templates(folder: Path, config: AppConfig) -> dict[str, TemplateDefinition]:
    if not folder.is_dir():
        raise TemplateError(f"模板目录不存在：{folder}")
    allowed = {
        _rgb_key(str(item))
        for item in config.app.get("templates", {}).get("selected_fill_rgb", [])
    }
    max_rows = int(
        config.app.get("templates", {}).get("max_header_scan_rows", 20)
    )
    keywords = config.test_types.get("template_keywords", {})
    definitions: dict[str, TemplateDefinition] = {}
    for path in sorted(folder.glob("*.xlsx")):
        test_type = _template_type(path, keywords)
        if not test_type:
            continue
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                workbook = load_workbook(
                    path, read_only=True, data_only=True, keep_links=False
                )
            sheet = workbook.active
            candidates: list[tuple[int, list[str], list[str]]] = []
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                for row_number, row in enumerate(
                    sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_rows)), 1
                ):
                    fields: list[str] = []
                    colors: list[str] = []
                    for cell in row:
                        rgb = _cell_rgb(cell)
                        if _rgb_key(rgb) in allowed and cell.value is not None:
                            name = str(cell.value).strip()
                            if name and name not in fields:
                                fields.append(name)
                                colors.append(rgb)
                    if fields:
                        candidates.append((row_number, fields, colors))
            sheet_name = sheet.title
            workbook.close()
        except Exception as exc:
            raise TemplateError(f"模板无法读取 {path.name}: {exc}") from exc
        if not candidates:
            raise TemplateError(f"模板没有黄色或绿色字段：{path.name}")
        header_row, fields, colors = max(candidates, key=lambda item: len(item[1]))
        if test_type in definitions:
            raise TemplateError(f"测试类型 {test_type} 匹配到多个模板")
        definitions[test_type] = TemplateDefinition(
            test_type=test_type,
            path=path,
            sheet_name=sheet_name,
            header_row=header_row,
            fields=fields,
            selected_colors=colors,
        )
    if not definitions:
        raise TemplateError("模板目录中没有找到可用的 xlsx 表头模板")
    return definitions
